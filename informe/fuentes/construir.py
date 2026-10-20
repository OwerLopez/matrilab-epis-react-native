"""Genera el informe del LAB 04 en LaTeX (main.tex + sections/) y en DOCX a partir de contenido.py.

Uso: python construir.py [--docx]
El PDF se compila después con Tectonic sobre ../main.tex.
"""
import sys
from pathlib import Path

from contenido import inf
from informe_lib import Docx, a_latex, tex

RAIZ = Path(__file__).resolve().parent.parent          # LAB 04/informe
CODIGO = Path("..")                                     # raíz del repositorio, relativa a main.tex
EQUIPO = ["Lopez Arela, Ower Frank", "Zegarra Puma, Mauricio Eduardo", "Huamani Chañi, Diego Oswaldo"]
DOCENTE = "M. Sc. Ing. R. Fabrizio Calienes Rodriguez"
TITULO = "Framework multiplataforma: interfaces de MatriLab EPIS con React Native"
FECHA = "20/10/2026"

PREAMBULO = r"""\documentclass[11pt, a4paper]{article}

% --- 1. GEOMETRÍA Y MÁRGENES ---
\usepackage[
    a4paper,
    left=2cm, right=2cm,
    top=1.2cm, bottom=2.0cm,
    includehead, headheight=105pt, headsep=0.5cm
]{geometry}

\usepackage[spanish,es-tabla,es-nodecimaldot]{babel}
\usepackage{fontspec}
\setmainfont{texgyretermes}[Extension=.otf, UprightFont=*-regular, BoldFont=*-bold, ItalicFont=*-italic, BoldItalicFont=*-bolditalic]
\setmonofont{DejaVuSansMono}[Extension=.ttf, BoldFont=*-Bold, ItalicFont=*-Oblique, BoldItalicFont=*-BoldOblique, Scale=0.84]
\usepackage{newunicodechar}
\newunicodechar{→}{\ensuremath{\rightarrow}}
\newunicodechar{−}{\ensuremath{-}}
\newunicodechar{≥}{\ensuremath{\geq}}
\newunicodechar{≤}{\ensuremath{\leq}}
\usepackage{microtype}
\usepackage{xurl}
\usepackage{graphicx}
\usepackage{fancyhdr}
\usepackage{xcolor}
\usepackage{colortbl}
\usepackage{tabularx}
\usepackage{longtable}
\usepackage{array}
\usepackage{setspace}
\usepackage{enumitem}
\usepackage{lastpage}
\usepackage{float}
\usepackage{needspace}
\usepackage[section]{placeins}
\usepackage{caption}
\usepackage{multirow}
\usepackage{amsmath}
\usepackage{amssymb}
\usepackage[
    colorlinks=true,
    urlcolor=colabpyblue,
    linkcolor=colabpyblue,
    citecolor=colabpyblue,
    breaklinks=true
]{hyperref}

% --- 2. LIBRERÍAS DE DISEÑO AVANZADO ---
\usepackage[most]{tcolorbox}
\tcbuselibrary{listings, skins, breakable}
\usepackage{fontawesome5}
\usepackage{pifont}
\usepackage{listings}

% --- 3. PALETA DE COLORES INSTITUCIONALES (UNSA) ---
\definecolor{unsared}{RGB}{230, 25, 45}
\definecolor{colabpyblue}{RGB}{25, 55, 90}
\definecolor{colabpybg}{RGB}{242, 247, 252}
\definecolor{colaboutgreen}{RGB}{45, 90, 60}
\definecolor{colaboutbg}{RGB}{245, 252, 245}
\definecolor{tableheader}{RGB}{242, 242, 242}
\definecolor{checkgreen}{RGB}{34, 139, 34}
\definecolor{codegray}{RGB}{248, 248, 248}
\definecolor{codecomment}{RGB}{110, 110, 110}
\definecolor{codestring}{RGB}{163, 21, 21}
\definecolor{codekeyword}{RGB}{25, 55, 160}

\emergencystretch=3em
\captionsetup{font=small, labelfont=bf, labelsep=period}
\renewcommand{\topfraction}{0.9}\renewcommand{\bottomfraction}{0.8}\renewcommand{\textfraction}{0.08}\renewcommand{\floatpagefraction}{0.75}\setcounter{topnumber}{3}\setcounter{totalnumber}{4}
\captionsetup[table]{position=top, skip=4pt}
\renewcommand{\listfigurename}{\textcolor{colabpyblue}{Índice de Figuras}}
\renewcommand{\listtablename}{\textcolor{colabpyblue}{Índice de Tablas}}
\setlength{\LTpre}{4pt}
\setlength{\LTpost}{8pt}

% --- 4. LISTADOS DE CÓDIGO ---
\lstdefinelanguage{TSX}{
    keywords={const, let, var, function, return, import, from, export, default, type, interface, if, else, switch, case,
              for, of, new, true, false, null, undefined, async, await, extends, typeof, in, as, throw, break, while},
    sensitive=true,
    comment=[l]{//},
    morecomment=[s]{/*}{*/},
    morestring=[b]', morestring=[b]", morestring=[b]`
}
\lstdefinelanguage{JSONx}{morestring=[b]", sensitive=true}
\lstdefinelanguage{Texto}{sensitive=true}
\lstset{
    basicstyle=\ttfamily\footnotesize,
    keywordstyle=\color{codekeyword}\bfseries,
    commentstyle=\color{codecomment}\itshape,
    stringstyle=\color{codestring},
    backgroundcolor=\color{codegray},
    frame=single, rulecolor=\color{colabpyblue!35},
    breaklines=true, breakatwhitespace=false, columns=fullflexible, keepspaces=true,
    showstringspaces=false, tabsize=2, upquote=true,
    numbers=none, xleftmargin=4pt, xrightmargin=4pt, aboveskip=6pt, belowskip=8pt,
    extendedchars=true,
    title={}, captionpos=t
}
\lstdefinestyle{anexo}{basicstyle=\ttfamily\scriptsize, numbers=left, numberstyle=\tiny\color{codecomment}, numbersep=6pt, xleftmargin=16pt}
\renewcommand{\lstlistingname}{Código}

% --- 5. COMPONENTES VISUALES Y MARCADORES ---
\newcommand{\cmark}{\textcolor{checkgreen}{\ding{51}}}
\newcommand{\xmark}{\textcolor{red!80!black}{\ding{55}}}

\newcommand{\safeimage}[3]{%
    \IfFileExists{#2}{%
        \includegraphics[width=#1, height=21cm, keepaspectratio]{#2}%
    }{%
        \fbox{\parbox{#1}{\centering\small [Imagen no encontrada: \texttt{#2}]}}%
    }%
}

\newcommand{\seccionInforme}[1]{
    \needspace{12\baselineskip}
    \begin{tcolorbox}[
        enhanced, sharp corners,
        colback=unsared, colframe=unsared,
        left=4pt, right=4pt, top=3pt, bottom=3pt,
        halign=center, before skip=16pt, after skip=10pt
    ]
        \hyphenpenalty=10000\exhyphenpenalty=10000\textcolor{white}{\textbf{\large\MakeUppercase{#1}}}
    \end{tcolorbox}
}

\newcommand{\colabtitle}[1]{
    \FloatBarrier
    \needspace{7\baselineskip}
    \vspace{0.4cm}
    {\large \textbf{\textcolor{colabpyblue}{#1}}}
    \vspace{0.1cm} \hrule height 1pt \relax \vspace{0.3cm}
}

\newcommand{\subsubtitulo}[1]{\FloatBarrier\par\needspace{9\baselineskip}\vspace{0.25cm}\noindent{\normalsize\textbf{\textcolor{unsared}{\faAngleRight}\ \textcolor{colabpyblue}{#1}}}\par\nopagebreak\vspace{0.15cm}}

\newtcolorbox{notaInfo}{enhanced, breakable, colback=colabpybg, colframe=colabpyblue, boxrule=0.6pt, leftrule=3pt, sharp corners, left=6pt, right=6pt, top=4pt, bottom=4pt, fontupper=\small}
\newtcolorbox{notaOk}{enhanced, breakable, colback=colaboutbg, colframe=colaboutgreen, boxrule=0.6pt, leftrule=3pt, sharp corners, left=6pt, right=6pt, top=4pt, bottom=4pt, fontupper=\small}
\newtcolorbox{notaAlerta}{enhanced, breakable, colback=unsared!6, colframe=unsared, boxrule=0.6pt, leftrule=3pt, sharp corners, left=6pt, right=6pt, top=4pt, bottom=4pt, fontupper=\small}

% --- 6. ENCABEZADO Y PIE DE PÁGINA OFICIAL EPIS-UNSA ---
\pagestyle{fancy}
\fancyhf{}
\setlength{\headwidth}{\textwidth}
\renewcommand{\headrulewidth}{0pt}
\fancyhead[C]{
    \fontsize{8.5}{10}\selectfont
    \begin{tabularx}{\linewidth}{|m{3.0cm}|X|m{2.5cm}|}
        \hline
        \centering\arraybackslash \includegraphics[width=2.5cm, height=1.6cm, keepaspectratio]{img/logo_episunsa.png} &
        \centering \textbf{UNIVERSIDAD NACIONAL DE SAN AGUSTÍN} \\
        FACULTAD DE INGENIERÍA DE PRODUCCIÓN Y SERVICIOS \\
        \textbf{ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS} &
        \centering\arraybackslash \includegraphics[width=1.5cm, height=1.5cm, keepaspectratio]{img/logo_abet.png} \tabularnewline
        \hline
        \multicolumn{3}{|c|}{\cellcolor{tableheader}\textbf{Formato: Guía de Práctica de Laboratorio / Talleres / Centros de Simulación}} \\
        \hline
        \centering\arraybackslash Aprobación: 2024/03/01 & \centering\arraybackslash Código: GUIA-PRLE-001 & \centering\arraybackslash Página: \thepage\ de \pageref{LastPage} \tabularnewline
        \hline
    \end{tabularx}
}
"""

PORTADA = r"""
\hypersetup{pdftitle={@TITULO@}, pdfauthor={@AUTORES@}, pdfsubject={Plataformas Emergentes (E) - Laboratorio 04}, pdfkeywords={React Native, Expo, multiplataforma, layout, flexbox}, pdfcreator={}, pdfproducer={}}
\begin{document}

\begin{center}
    \vspace*{0.2cm}
    {\LARGE \textbf{\textcolor{colabpyblue}{INFORME DE LABORATORIO}}} \\[0.15cm]
    {\large \textbf{Plataformas Emergentes (E) -- EPIS UNSA}}
\end{center}
\vspace{0.25cm}

\noindent
\renewcommand{\arraystretch}{1.38}
\setlength{\tabcolsep}{6pt}
\begin{tabularx}{\linewidth}{|>{\bfseries\columncolor{tableheader}\raggedright\arraybackslash}p{3.6cm}|X|>{\bfseries\columncolor{tableheader}\raggedright\arraybackslash}p{3.2cm}|X|}
\hline
\rowcolor{unsared}
\multicolumn{4}{|c|}{\textcolor{white}{\Large\textbf{I. INFORMACIÓN BÁSICA}}} \\
\hline
ASIGNATURA: &
\multicolumn{3}{|>{\raggedright\arraybackslash}p{\dimexpr\linewidth-3.6cm-4\tabcolsep-2\arrayrulewidth\relax}|}{%
    \textbf{PLATAFORMAS EMERGENTES (E)}%
} \\
\hline
TÍTULO DE PRÁCTICA: &
\multicolumn{3}{|>{\raggedright\arraybackslash}p{\dimexpr\linewidth-3.6cm-4\tabcolsep-2\arrayrulewidth\relax}|}{%
    @TITULO_TEX@%
} \\
\hline
NRO. PRÁCTICA: & 04 & AÑO LECTIVO: & 2026-B \\
\hline
SEMESTRE ACADÉMICO: & VIII & FECHA DE ENTREGA: & @FECHA@ \\
\hline
MODALIDAD / TIPO: & Grupal (3 integrantes) & CURSO: & LABORATORIO \\
\hline
ESTUDIANTE(S): &
\multicolumn{3}{|>{\raggedright\arraybackslash}p{\dimexpr\linewidth-3.6cm-4\tabcolsep-2\arrayrulewidth\relax}|}{%
    @ESTUDIANTES@%
} \\
\hline
DOCENTE: &
\multicolumn{3}{|>{\raggedright\arraybackslash}p{\dimexpr\linewidth-3.6cm-4\tabcolsep-2\arrayrulewidth\relax}|}{%
    @DOCENTE@%
} \\
\hline
REPOSITORIO: &
\multicolumn{3}{|>{\raggedright\arraybackslash}p{\dimexpr\linewidth-3.6cm-4\tabcolsep-2\arrayrulewidth\relax}|}{%
    \url{https://github.com/OwerLopez/matrilab-epis-react-native}%
} \\
\hline
\end{tabularx}
\vspace{0.6cm}

\begin{center}
\includegraphics[width=0.78\textwidth]{img/figuras/fig_i04.png}\\[0.2cm]
{\small\itshape MatriLab EPIS ejecutándose en Android: mockup, grupo disponible, bloqueo por cruce de horario y grupo agotado.}
\end{center}
\clearpage

\renewcommand{\contentsname}{\textcolor{colabpyblue}{Índice General de Contenidos}}
{
  \hypersetup{linkcolor=colabpyblue}
  \setcounter{tocdepth}{2}
  \small
  \tableofcontents
}
\clearpage
{\hypersetup{linkcolor=colabpyblue}\listoffigures}
\vspace{0.6cm}
{\hypersetup{linkcolor=colabpyblue}\listoftables}
\clearpage
"""


def main():
    secciones = a_latex(inf, RAIZ / "sections", CODIGO)
    estudiantes = r" \newline ".join(f"\\textbf{{{i}.}} {tex(n)}" for i, n in enumerate(EQUIPO, 1))
    portada = (PORTADA.replace("@TITULO@", TITULO).replace("@AUTORES@", "; ".join(EQUIPO))
               .replace("@TITULO_TEX@", tex(TITULO)).replace("@FECHA@", FECHA)
               .replace("@ESTUDIANTES@", estudiantes).replace("@DOCENTE@", DOCENTE))
    cuerpo = "".join(f"\\input{{sections/{arch}}}\n" for _, arch in secciones)
    (RAIZ / "main.tex").write_text(PREAMBULO + portada + cuerpo + "\n\\end{document}\n", encoding="utf-8", newline="\n")
    print("main.tex con", len(secciones), "secciones")
    if "--docx" not in sys.argv:
        return

    d = Docx()
    d.encabezado(RAIZ / "img" / "logo_episunsa.png", RAIZ / "img" / "logo_abet.png")
    d.portada([
        ("ASIGNATURA:", "**PLATAFORMAS EMERGENTES (E)**"),
        ("TÍTULO DE PRÁCTICA:", TITULO),
        ("NRO. PRÁCTICA:", "04", "AÑO LECTIVO:", "2026-B"),
        ("SEMESTRE ACADÉMICO:", "VIII", "FECHA DE ENTREGA:", FECHA),
        ("MODALIDAD / TIPO:", "Grupal (3 integrantes)", "CURSO:", "LABORATORIO"),
        ("ESTUDIANTE(S):", "\n".join(f"**{i}.** {n}" for i, n in enumerate(EQUIPO, 1))),
        ("DOCENTE:", DOCENTE),
        ("REPOSITORIO:", "https://github.com/OwerLopez/matrilab-epis-react-native"),
    ], "INFORME DE LABORATORIO", "Plataformas Emergentes (E) – EPIS UNSA")
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm
    p = d.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(RAIZ / "img" / "figuras" / "fig_i04.png"), width=Cm(11.5))
    p = d.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MatriLab EPIS ejecutándose en Android: mockup, grupo disponible, bloqueo por cruce de horario y grupo agotado."); r.italic = True
    d.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.indices()
    d.construir(inf, RAIZ, RAIZ.parent, ("Lopez Arela, Ower Frank", "OL"))
    d.metadatos(TITULO, "; ".join(EQUIPO), "Plataformas Emergentes (E) – Laboratorio 04", "React Native; Expo; multiplataforma; layout; flexbox")
    salida = RAIZ / "LAB04_Plataformas_Emergentes_MatriLab_React_Native.docx"
    d.guardar(salida)
    print("docx", salida.name)


if __name__ == "__main__":
    main()
