"""Bloques del informe y sus dos salidas: secciones LaTeX (formato EPIS de los laboratorios) y DOCX.

El contenido se escribe una sola vez con marcas ligeras en línea:
  **negrita**, *cursiva*, `código` y <https://url>.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROJO = RGBColor(0xE6, 0x19, 0x2D)
AZUL = RGBColor(0x19, 0x37, 0x5A)
GRIS = RGBColor(0x55, 0x55, 0x55)
MARCAS = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|<https?://[^>]+>)")


# ------------------------------------------------------------------ modelo
class Informe:
    def __init__(self):
        self.bloques = []

    def _b(self, *b):
        self.bloques.append(b)

    def seccion(self, t): self._b("sec", t)
    def sub(self, t): self._b("sub", t)
    def subsub(self, t, nueva_pagina=False): self._b("subsub", t, nueva_pagina)
    def p(self, t): self._b("p", t)
    def lista(self, items, numerada=False): self._b("lista", items, numerada)
    def figura(self, ruta, leyenda, ancho=0.95, etiqueta=None): self._b("fig", ruta, leyenda, ancho, etiqueta)
    def tabla(self, leyenda, cabecera, filas, anchos, tam="small"): self._b("tab", leyenda, cabecera, filas, anchos, tam)
    def codigo(self, texto, titulo=None, leng="TSX"): self._b("cod", texto.strip("\n"), titulo, leng)
    def archivo(self, ruta, titulo): self._b("arch", ruta, titulo)
    def nota(self, t, tipo="info", comentario=None): self._b("nota", t, tipo, comentario)
    def salto(self): self._b("salto",)
    def referencias(self, items): self._b("refs", items)


# ------------------------------------------------------------------ LaTeX
ESC = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
       "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


def esc(t):
    return "".join(ESC.get(ch, ch) for ch in t)


def esc_tt(t):
    # En \texttt se permite cortar la línea después de puntos, barras, guiones bajos y signos igual,
    # y entre una minúscula y una mayúscula de los identificadores camelCase.
    s = esc(t)
    s = re.sub(r"(\\_|/|\.|=)", r"\1\\allowbreak{}", s)
    return re.sub(r"([a-z])([A-Z])", r"\1\\allowbreak{}\2", s)


def tex(t):
    out = []
    for parte in MARCAS.split(t):
        if not parte:
            continue
        if parte.startswith("**"):
            out.append(r"\textbf{" + tex(parte[2:-2]) + "}")
        elif parte.startswith("`"):
            out.append(r"\texttt{" + esc_tt(parte[1:-1]) + "}")
        elif parte.startswith("<http"):
            out.append(r"\url{" + parte[1:-1].replace("%", r"\%").replace("#", r"\#") + "}")
        elif parte.startswith("*"):
            out.append(r"\textit{" + tex(parte[1:-1]) + "}")
        else:
            out.append(esc(parte))
    return "".join(out)


def a_latex(inf, carpeta, raiz_codigo):
    """Escribe una sección por archivo y devuelve la lista de (título, archivo) para main.tex."""
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    archivos, actual, n = [], None, 0

    def nuevo(titulo):
        nonlocal actual, n
        n += 1
        slug = re.sub(r"[^a-z0-9]+", "_", titulo.lower().translate(str.maketrans("áéíóúñ", "aeioun"))).strip("_")
        nombre = f"{n:02d}_{slug[:40]}.tex"
        actual = open(carpeta / nombre, "w", encoding="utf-8", newline="\n")
        archivos.append((titulo, nombre))
        return actual

    for b in inf.bloques:
        k = b[0]
        if k == "sec":
            if actual:
                actual.close()
            f = nuevo(b[1])
            f.write(("\\clearpage\n" if n > 1 else "") + f"\\seccionInforme{{{tex(b[1])}}}\n\\phantomsection\\addcontentsline{{toc}}{{section}}{{{tex(b[1])}}}\n\n")
            continue
        f = actual
        if k == "sub":
            f.write(f"\\vspace{{0.2cm}}\n\\phantomsection\\addcontentsline{{toc}}{{subsection}}{{{tex(b[1])}}}\n\\colabtitle{{{tex(b[1])}}}\n\n")
        elif k == "subsub":
            f.write(("\\clearpage\n" if len(b) > 2 and b[2] else "") + f"\\subsubtitulo{{{tex(b[1])}}}\n\n")
        elif k == "p":
            f.write(tex(b[1]) + "\n\n")
        elif k == "lista":
            entorno = "enumerate" if b[2] else "itemize"
            opc = "[label=\\textbf{\\arabic*.}, leftmargin=*]" if b[2] else "[label=\\textcolor{unsared}{\\textbullet}, leftmargin=*]"
            f.write(f"\\begin{{{entorno}}}{opc}\n" + "".join(f"    \\item {tex(i)}\n" for i in b[1]) + f"\\end{{{entorno}}}\n\n")
        elif k == "fig":
            _, ruta, ley, ancho, et = b
            f.write("\\begin{figure}[H]\n\\centering\n"
                    f"\\safeimage{{{ancho}\\textwidth}}{{{ruta}}}{{{esc(ley)}}}\n"
                    f"\\caption{{{tex(ley)}}}\n" + (f"\\label{{{et}}}\n" if et else "") + "\\end{figure}\n\n")
        elif k == "tab":
            _, ley, cab, filas, anchos, tam = b
            cols = "|" + "|".join(f">{{\\raggedright\\arraybackslash}}p{{\\dimexpr {a:.3f}\\linewidth-2\\tabcolsep-2\\arrayrulewidth\\relax}}" for a in anchos) + "|"
            enc = " & ".join(f"\\color{{white}}\\textbf{{{tex(c)}}}" for c in cab)
            # Todas las tablas del informe caben en una página: se componen como un bloque indivisible
            # con su leyenda arriba, igual que una tabla flotante [H].
            f.write(f"\\begin{{table}}[H]\n\\centering\\{tam}\n\\caption{{{tex(ley)}}}\n\\renewcommand{{\\arraystretch}}{{1.25}}\n"
                    f"\\begin{{tabular}}{{{cols}}}\n\\hline\n\\rowcolor{{colabpyblue}} {enc} \\\\ \\hline\n")
            for i, fila in enumerate(filas):
                color = "\\rowcolor{colabpybg} " if i % 2 else ""
                f.write(color + " & ".join(tex(str(c)) for c in fila) + " \\\\ \\hline\n")
            f.write("\\end{tabular}\n\\end{table}\n\n")
        elif k == "cod":
            _, texto, titulo, leng = b
            t = f", title={{{esc(titulo)}}}" if titulo else ""
            lineas = texto.count("\n") + 1
            if lineas <= 14:
                f.write(f"\\needspace{{{lineas + 3}\\baselineskip}}\n")
            f.write(f"\\begin{{lstlisting}}[language={leng}{t}]\n{texto}\n\\end{{lstlisting}}\n\\par\\penalty-100\n\n")
        elif k == "arch":
            _, ruta, titulo = b
            leng = {".json": "JSONx", ".mjs": "TSX", ".ts": "TSX", ".tsx": "TSX"}.get(Path(ruta).suffix, "TSX")
            f.write(f"\\subsubtitulo{{{tex(titulo)}}}\n\\lstinputlisting[language={leng}, style=anexo, title={{\\texttt{{{esc(ruta)}}}}}]{{{(Path(raiz_codigo) / ruta).as_posix()}}}\n\n")
        elif k == "nota":
            _, texto, tipo, _c = b
            caja = {"info": "notaInfo", "ok": "notaOk", "alerta": "notaAlerta"}[tipo]
            f.write(f"\\begin{{{caja}}}\n{tex(texto)}\n\\end{{{caja}}}\n\n")
        elif k == "salto":
            f.write("\\clearpage\n\n")
        elif k == "refs":
            f.write("\\begin{enumerate}[label={[\\arabic*]}, leftmargin=*, itemsep=3pt]\n"
                    + "".join(f"    \\item {tex(i)}\n" for i in b[1]) + "\\end{enumerate}\n\n")
    if actual:
        actual.close()
    return archivos


# ------------------------------------------------------------------ DOCX
ORDEN_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
               "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def _ordenar_tblpr(tabla):
    """Word valida el orden de los hijos de w:tblPr; se reordenan según el esquema OOXML."""
    tblPr = tabla._tbl.tblPr
    hijos = list(tblPr)
    clave = lambda e: ORDEN_TBLPR.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDEN_TBLPR else 99
    for e in hijos:
        tblPr.remove(e)
    for e in sorted(hijos, key=clave):
        tblPr.append(e)


def _sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _bordes(tabla, color="9A9A9A", sz="4"):
    tblPr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for borde in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{borde}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)
    _ordenar_tblpr(tabla)


def _anchos(tabla, cm):
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    for viejo in tblPr.findall(qn("w:tblLayout")) + tblPr.findall(qn("w:tblW")):
        tblPr.remove(viejo)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(int(sum(cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    for col, a in zip(tabla._tbl.tblGrid.findall(qn("w:gridCol")), cm):
        col.set(qn("w:w"), str(int(a * 567)))
    for fila in tabla.rows:
        for celda, a in zip(fila.cells, cm):
            celda.width = Cm(a)
    _ordenar_tblpr(tabla)


def _cabecera_repetida(fila):
    trPr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def campo(parrafo, instruccion, texto="1"):
    r = parrafo.add_run()
    for tipo in ("begin", "instr", "separate", "text", "end"):
        if tipo == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instruccion
        elif tipo == "text":
            el = OxmlElement("w:t")
            el.text = texto
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        r._r.append(el)
    return r


class Docx:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = "Times New Roman"
        n.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        n.font.size = Pt(11)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.1
        for nivel, tam in ((1, 13), (2, 12), (3, 11)):
            h = st[f"Heading {nivel}"]
            h.font.name = "Times New Roman"
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            h.font.size = Pt(tam)
            h.font.bold = True
            h.font.color.rgb = AZUL if nivel > 1 else RGBColor(0xFF, 0xFF, 0xFF)
            h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        h1 = st["Heading 1"]
        h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pPr = h1.element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "E6192D")
        pPr.insert_element_before(shd, "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:contextualSpacing", "w:jc", "w:outlineLvl", "w:rPr")
        cap = st["Caption"]
        cap.font.name = "Times New Roman"
        cap.font.size = Pt(9.5)
        cap.font.italic = False
        cap.font.color.rgb = GRIS
        for nombre in ("toc 1", "toc 2", "table of figures"):
            try:
                e = st[nombre]
            except KeyError:
                from docx.enum.style import WD_STYLE_TYPE
                e = st.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
            e.base_style = st["Normal"]
            e.font.size = Pt(10)
            e.paragraph_format.space_after = Pt(1)
            e.paragraph_format.left_indent = Cm(0.6 if nombre == "toc 2" else 0)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(2.0)
        sec.header_distance = Cm(0.8)
        self.ancho_util = 17.0

    # ------- texto en línea
    def runs(self, p, texto, tam=None, color=None, negrita=False):
        primero = None
        for parte in MARCAS.split(texto):
            if not parte:
                continue
            if parte.startswith("**"):
                r = p.add_run(parte[2:-2]); r.bold = True
            elif parte.startswith("`"):
                r = p.add_run(parte[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((tam or 11) - 1.5)
            elif parte.startswith("<http"):
                r = p.add_run(parte[1:-1]); r.font.color.rgb = AZUL; r.underline = True
            elif parte.startswith("*"):
                r = p.add_run(parte[1:-1]); r.italic = True
            else:
                r = p.add_run(parte)
            if negrita:
                r.bold = True
            if tam and not parte.startswith("`"):
                r.font.size = Pt(tam)
            if color is not None:
                r.font.color.rgb = color
            primero = primero or r
        return primero

    def encabezado(self, logo_epis, logo_abet):
        """Réplica del encabezado institucional del formato GUIA-PRLE-001."""
        cab = self.doc.sections[0].header
        cab.is_linked_to_previous = False
        t = cab.add_table(rows=3, cols=3, width=Cm(17))
        _bordes(t, "000000", "6")
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells
        c[0].paragraphs[0].add_run().add_picture(str(logo_epis), width=Cm(2.4))
        c[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        par = c[1].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, (txt, b) in enumerate((("UNIVERSIDAD NACIONAL DE SAN AGUSTÍN", True), ("FACULTAD DE INGENIERÍA DE PRODUCCIÓN Y SERVICIOS", False), ("ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS", True))):
            r = par.add_run(("\n" if i else "") + txt); r.bold = b; r.font.size = Pt(8.5)
        c[2].paragraphs[0].add_run().add_picture(str(logo_abet), width=Cm(1.4))
        c[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        fila = t.rows[1].cells
        m = fila[0].merge(fila[1]).merge(fila[2])
        m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = m.paragraphs[0].add_run("Formato: Guía de Práctica de Laboratorio / Talleres / Centros de Simulación"); r.bold = True; r.font.size = Pt(8.5)
        _sombrear(m, "F2F2F2")
        f = t.rows[2].cells
        for celda, txt in zip(f, ("Aprobación: 2024/03/01", "Código: GUIA-PRLE-001", None)):
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if txt:
                p.add_run(txt).font.size = Pt(8.5)
            else:
                p.add_run("Página: ").font.size = Pt(8.5)
                campo(p, "PAGE").font.size = Pt(8.5)
                p.add_run(" de ").font.size = Pt(8.5)
                campo(p, "NUMPAGES").font.size = Pt(8.5)
        for row in t.rows:
            for celda in row.cells:
                for p in celda.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
        _anchos(t, [4.0, 9.6, 3.4])
        cab.paragraphs[0].paragraph_format.space_after = Pt(0)

    def portada(self, filas, titulo, subtitulo):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.runs(p, titulo, tam=18, color=AZUL, negrita=True)
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.runs(p, subtitulo, tam=13, negrita=True)
        t = self.doc.add_table(rows=1, cols=4)
        _bordes(t, "000000", "6")
        m = t.rows[0].cells[0].merge(t.rows[0].cells[3])
        _sombrear(m, "E6192D")
        m.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = m.paragraphs[0].add_run("I. INFORMACIÓN BÁSICA"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for fila in filas:
            celdas = t.add_row().cells
            if len(fila) == 2:
                destino = celdas[1].merge(celdas[3])
                pares = [(celdas[0], fila[0], True), (destino, fila[1], False)]
            else:
                pares = [(celdas[i], fila[i], i % 2 == 0) for i in range(4)]
            for celda, txt, etiqueta in pares:
                celda.paragraphs[0].text = ""
                for j, linea in enumerate(txt.split("\n")):
                    par = celda.paragraphs[0] if j == 0 else celda.add_paragraph()
                    par.paragraph_format.space_after = Pt(1)
                    self.runs(par, linea, tam=10.5, negrita=etiqueta)
                if etiqueta:
                    _sombrear(celda, "F2F2F2")
        _anchos(t, [3.8, 4.7, 3.6, 4.9])

    def indices(self):
        for titulo, instr in (("Índice general de contenidos", 'TOC \\o "1-2" \\h \\z \\u'),
                              ("Índice de figuras", 'TOC \\h \\z \\c "Figura"'),
                              ("Índice de tablas", 'TOC \\h \\z \\c "Tabla"')):
            p = self.doc.add_paragraph()
            self.runs(p, titulo, tam=14, color=AZUL, negrita=True)
            campo(self.doc.add_paragraph(), instr, "Actualice los campos para ver el índice.")
            self.doc.add_paragraph()

    def leyenda(self, tipo, texto, antes=False):
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = antes
        r = p.add_run(f"{tipo} "); r.bold = True
        campo(p, f"SEQ {tipo} \\* ARABIC").bold = True
        p.add_run(". ").bold = True
        self.runs(p, texto)

    def construir(self, inf, raiz_img, raiz_codigo, autor):
        doc = self.doc
        for b in inf.bloques:
            k = b[0]
            if k == "sec":
                h = doc.add_heading(b[1].upper(), level=1)
                h.paragraph_format.page_break_before = True
            elif k == "sub":
                doc.add_heading(b[1], level=2)
            elif k == "subsub":
                h = doc.add_heading(b[1], level=3)
                if len(b) > 2 and b[2]:
                    h.paragraph_format.page_break_before = True
            elif k == "p":
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                self.runs(p, b[1])
            elif k == "lista":
                for i in b[1]:
                    p = doc.add_paragraph(style="List Number" if b[2] else "List Bullet")
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p.paragraph_format.space_after = Pt(3)
                    self.runs(p, i)
            elif k == "fig":
                _, ruta, ley, ancho, _e = b
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                from PIL import Image
                with Image.open(Path(raiz_img) / ruta) as im:
                    prop = im.height / im.width
                w = min(self.ancho_util * ancho, 21.5 / prop)
                p.add_run().add_picture(str(Path(raiz_img) / ruta), width=Cm(w))
                self.leyenda("Figura", ley)
            elif k == "tab":
                _, ley, cab, filas, anchos, tam = b
                self.leyenda("Tabla", ley, antes=True)
                t = doc.add_table(rows=1, cols=len(cab))
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                _bordes(t)
                pt = 9 if tam == "small" else 8.5 if tam == "footnotesize" else 10
                for i, c in enumerate(cab):
                    celda = t.rows[0].cells[i]
                    _sombrear(celda, "19375A")
                    celda.paragraphs[0].paragraph_format.space_after = Pt(0)
                    self.runs(celda.paragraphs[0], c, tam=pt, color=RGBColor(0xFF, 0xFF, 0xFF), negrita=True)
                _cabecera_repetida(t.rows[0])
                for j, fila in enumerate(filas):
                    celdas = t.add_row().cells
                    for i, v in enumerate(fila):
                        par = celdas[i].paragraphs[0]
                        par.paragraph_format.space_after = Pt(0)
                        par.paragraph_format.line_spacing = 1.0
                        self.runs(par, str(v), tam=pt)
                        if j % 2:
                            _sombrear(celdas[i], "F2F7FC")
                _anchos(t, [a * self.ancho_util for a in anchos])
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif k in ("cod", "arch"):
                if k == "cod":
                    _, texto, titulo, _l = b
                else:
                    _, ruta, titulo = b
                    doc.add_heading(titulo, level=3)
                    texto = (Path(raiz_codigo) / ruta).read_text(encoding="utf-8").rstrip("\n")
                    titulo = ruta
                self.bloque_codigo(texto, titulo, 7.5 if k == "arch" else 8.5)
            elif k == "nota":
                _, texto, tipo, comentario = b
                t = doc.add_table(rows=1, cols=1)
                c = t.cell(0, 0)
                fondo, borde = {"info": ("F2F7FC", "19375A"), "ok": ("F5FCF5", "2D5A3C"), "alerta": ("FDECEE", "E6192D")}[tipo]
                _sombrear(c, fondo)
                _bordes(t, borde, "8")
                _anchos(t, [self.ancho_util])
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r = self.runs(c.paragraphs[0], texto, tam=10)
                if comentario:
                    doc.add_comment(r, text=comentario, author=autor[0], initials=autor[1])
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif k == "salto":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif k == "refs":
                for i, ref in enumerate(b[1], 1):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.9)
                    p.paragraph_format.first_line_indent = Cm(-0.9)
                    p.paragraph_format.space_after = Pt(4)
                    p.add_run(f"[{i}]	")
                    self.runs(p, ref, tam=10.5)

    def bloque_codigo(self, texto, titulo, tam):
        t = self.doc.add_table(rows=1, cols=1)
        c = t.cell(0, 0)
        _sombrear(c, "F8F8F8")
        _bordes(t, "C8C8C8")
        _anchos(t, [self.ancho_util])
        c.paragraphs[0].text = ""
        primera = True
        if titulo:
            p = c.paragraphs[0]
            r = p.add_run(titulo); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = AZUL
            p.paragraph_format.space_after = Pt(3)
            primera = False
        for linea in texto.split("\n"):
            p = c.paragraphs[0] if primera else c.add_paragraph()
            primera = False
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(linea.replace("\t", "  "))
            r.font.name = "Consolas"
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.size = Pt(tam)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def metadatos(self, titulo, autores, asunto, claves):
        cp = self.doc.core_properties
        cp.title, cp.author, cp.last_modified_by = titulo, autores, autores.split(";")[0].strip()
        cp.subject, cp.keywords, cp.comments, cp.category = asunto, claves, "", "Informe de laboratorio"
        cp.language = "es-PE"
        cp.revision = 1
        import datetime
        cp.created = cp.modified = datetime.datetime.now().replace(microsecond=0)

    def guardar(self, ruta):
        self.doc.save(str(ruta))
